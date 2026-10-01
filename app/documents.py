"""Bounded file extraction and fresh, text-focused document exports."""
import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches as WordInches, Pt as WordPt, RGBColor as WordColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

from .pricing import count_words


class DocumentError(ValueError):
    pass


def clean_text(text):
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    # Preserve language-specific joiners, punctuation, citations and meaningful Unicode.
    return "".join(c for c in text if c in "\n\t" or ord(c) >= 32).strip()


def block(kind, text, **extra):
    return {"type": kind, "text": clean_text(text), **extra}


def normalize_blocks(blocks):
    result = []
    for item in blocks:
        if not item.get("text", "").strip():
            continue
        # Keep each provider request bounded, including giant unbroken paragraphs.
        text = item["text"]
        while text:
            cut = min(len(text), 4500)
            if cut < len(text):
                space = text.rfind(" ", 0, cut)
                if space > 3000:
                    cut = space
            part, text = text[:cut].strip(), text[cut:].strip()
            if part:
                result.append({**item, "id": f"b{len(result):04d}", "text": part})
    if len(result) > 900:
        raise DocumentError("This document has too many separate text blocks. Please split it into smaller files.")
    return result


def validate_zip(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 5000 or sum(i.file_size for i in infos) > 80 * 1024 * 1024:
                raise DocumentError("This document expands beyond the safe processing limit.")
            if any(i.flag_bits & 1 for i in infos):
                raise DocumentError("Password-protected files are not supported.")
            if any("vbaproject" in i.filename.lower() for i in infos):
                raise DocumentError("Please upload a document without macros.")
    except zipfile.BadZipFile as exc:
        raise DocumentError("This is not a valid Office document.") from exc


def extract_docx(data):
    validate_zip(data)
    doc = WordDocument(io.BytesIO(data))
    blocks = []
    from docx.table import Table
    from docx.text.paragraph import Paragraph as WordParagraph
    for child in doc.element.body:
        if child.tag.endswith("}p"):
            p = WordParagraph(child, doc)
            kind = "heading" if p.style and p.style.name.startswith(("Heading", "Title")) else "paragraph"
            blocks.append(block(kind, p.text))
        elif child.tag.endswith("}tbl"):
            table = Table(child, doc)
            for row in table.rows:
                blocks.append(block("table_row", " | ".join(c.text for c in row.cells)))
    return blocks


def extract_pptx(data):
    validate_zip(data)
    prs = Presentation(io.BytesIO(data))
    if len(prs.slides) > 80:
        raise DocumentError("Please upload at most 80 slides at a time.")
    blocks = []
    for index, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.has_text_frame:
                kind = "heading" if shape == slide.shapes.title else "paragraph"
                for p in shape.text_frame.paragraphs:
                    blocks.append(block(kind, p.text, slide=index))
            if shape.has_table:
                for row in shape.table.rows:
                    blocks.append(block("table_row", " | ".join(c.text for c in row.cells), slide=index))
    return blocks, len(prs.slides)


def extract_pdf(data, ocr_language="eng"):
    if not data.lstrip().startswith(b"%PDF-"):
        raise DocumentError("This file does not look like a PDF.")
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise DocumentError("Please unlock this PDF before uploading it.")
    if len(reader.pages) > 80:
        raise DocumentError("Please upload at most 80 PDF pages at a time.")
    blocks, warnings, ocr_pages = [], [], 0
    ocr_ready = bool(shutil.which("pdftoppm") and shutil.which("tesseract"))
    with tempfile.TemporaryDirectory(prefix="txtzi-pdf-") as folder:
        source = Path(folder) / "source.pdf"
        for index, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if len(text.strip()) < 30:
                if not ocr_ready:
                    raise DocumentError("This PDF contains scanned or unreadable pages. OCR is not installed on this server. Upload a text-based PDF or paste the text.")
                if ocr_pages >= 25:
                    raise DocumentError("Please split scanned PDFs into files of at most 25 scanned pages.")
                if not source.exists():
                    source.write_bytes(data)
                image = Path(folder) / "page"
                subprocess.run(["pdftoppm", "-f", str(index), "-l", str(index), "-singlefile", "-scale-to", "1800", "-png", str(source), str(image)], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=35)
                result = subprocess.run(["tesseract", str(image) + ".png", "stdout", "-l", ocr_language], capture_output=True, check=True, timeout=35)
                text = result.stdout.decode("utf-8", errors="replace")
                ocr_pages += 1
            for paragraph in re.split(r"\n\s*\n", text):
                blocks.append(block("paragraph", paragraph, page=index))
    if ocr_pages:
        warnings.append(f"Text was read from {ocr_pages} scanned page(s). Review names, numbers and punctuation for OCR errors.")
    return blocks, ocr_pages, warnings


def extract_file(filename, data, language="English"):
    extension = Path(filename).suffix.lower()
    warnings = []
    ocr_pages, slides = 0, 0
    try:
        if extension == ".pdf":
            lang = {"Hebrew": "heb+eng", "Arabic": "ara+eng", "Spanish": "spa+eng", "French": "fra+eng", "German": "deu+eng"}.get(language, "eng")
            blocks, ocr_pages, warnings = extract_pdf(data, lang)
        elif extension == ".docx":
            blocks = extract_docx(data)
        elif extension == ".pptx":
            blocks, slides = extract_pptx(data)
        elif extension == ".doc":
            if not data.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                raise DocumentError("Please save this document as DOCX and try again.")
            if not shutil.which("antiword"):
                raise DocumentError("Legacy DOC conversion is unavailable here. Please save as DOCX.")
            with tempfile.TemporaryDirectory(prefix="txtzi-doc-") as folder:
                path = Path(folder) / "input.doc"
                path.write_bytes(data)
                result = subprocess.run(["antiword", "-m", "UTF-8.txt", str(path)], capture_output=True, check=True, timeout=30)
                blocks = [block("paragraph", t) for t in result.stdout.decode("utf-8", errors="replace").split("\n\n")]
        elif extension in (".txt", ".md"):
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raise DocumentError("Please save text files as UTF-8.")
            blocks = [block("paragraph", t) for t in text.split("\n\n")]
        else:
            raise DocumentError("Choose PDF, DOC, DOCX, PPTX, TXT or Markdown.")
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("We could not read this file. Try exporting it again or paste its text.") from exc
    if extension in (".pdf", ".doc", ".docx", ".pptx"):
        warnings.append("Exports use a fresh, text-focused layout. Original images, charts, signatures, comments and exact page styling are not carried over.")
    result = normalize_blocks(blocks)
    return {"blocks": result, "word_count": count_words("\n".join(b["text"] for b in result)), "ocr_pages": ocr_pages, "slides": slides, "warnings": warnings, "source_type": extension[1:]}


def is_rtl(text):
    return len(re.findall(r"[\u0590-\u08ff]", text)) > len(re.findall(r"[A-Za-z]", text))


def export_body(title, blocks):
    if blocks and blocks[0]["type"] == "heading" and blocks[0]["text"].strip().casefold() == title.strip().casefold():
        return blocks[1:]
    return blocks


def export_docx(title, blocks, clean_metadata=True):
    blocks = export_body(title, blocks)
    doc = WordDocument()
    section = doc.sections[0]
    section.page_width = WordInches(8.5)
    section.page_height = WordInches(11)
    section.top_margin = WordInches(0.85)
    section.bottom_margin = WordInches(0.85)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = WordPt(11)
    normal.paragraph_format.space_after = WordPt(9)
    normal.paragraph_format.line_spacing = 1.2
    for name in ("Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.color.rgb = WordColor(0, 0, 0)
        style.font.underline = False
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    doc.add_paragraph(title, style="Title")
    for b in blocks:
        p = doc.add_heading(b["text"], 2) if b["type"] == "heading" else doc.add_paragraph(b["text"])
        if is_rtl(b["text"]):
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    props = doc.core_properties
    props.title = title
    props.author = "" if clean_metadata else "txtzi"
    props.last_modified_by = "" if clean_metadata else "txtzi"
    props.comments = ""
    props.keywords = ""
    props.subject = ""
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def font_paths():
    roots = [Path("/usr/share/fonts/truetype/dejavu"), Path("/usr/share/fonts/truetype/noto"), Path("/opt/codex/runtimes/codex-primary-runtime/dependencies/share/fonts")]
    for root in roots:
        path = root / "DejaVuSans.ttf"
        if path.exists():
            return path, root / "DejaVuSans-Bold.ttf"
    return None, None


def export_pdf(title, blocks, clean_metadata=True):
    blocks = export_body(title, blocks)
    regular, bold = font_paths()
    if regular:
        if "txtziSans" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("txtziSans", str(regular)))
            pdfmetrics.registerFont(TTFont("txtziBold", str(bold if bold.exists() else regular)))
        font, boldfont = "txtziSans", "txtziBold"
    else:
        font, boldfont = "Helvetica", "Helvetica-Bold"
    out = io.BytesIO()
    document = SimpleDocTemplate(out, pagesize=(595.28, 841.89), leftMargin=54, rightMargin=54, topMargin=54, bottomMargin=54, title=title, author="" if clean_metadata else "txtzi")
    story = []
    for b in [{"type": "title", "text": title}, *blocks]:
        text = b["text"]
        rtl = is_rtl(text)
        if rtl:
            try:
                from bidi.algorithm import get_display
                import arabic_reshaper
                text = get_display(arabic_reshaper.reshape(text))
            except ImportError:
                pass
        heading = b["type"] in ("heading", "title")
        style = ParagraphStyle("body", fontName=boldfont if heading else font, fontSize=22 if b["type"] == "title" else 13 if heading else 10.5, leading=29 if b["type"] == "title" else 17, textColor=colors.HexColor("#222827"), spaceAfter=16 if heading else 10, alignment=TA_RIGHT if rtl else TA_LEFT, wordWrap="CJK")
        story.append(Paragraph(escape(text).replace("\n", "<br/>"), style))
    def footer(canvas, doc):
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.HexColor("#7a827f"))
        canvas.drawRightString(541, 30, str(doc.page))
    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def export_pptx(title, blocks, clean_metadata=True):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    groups = []
    current = []
    last_slide = None
    for b in blocks:
        slide = b.get("slide")
        if current and ((slide and slide != last_slide) or sum(len(i["text"]) for i in current) + len(b["text"]) > 1100):
            groups.append(current)
            current = []
        # A long block is split into small chunks to avoid text overflow.
        remaining = b["text"]
        while remaining:
            cut = min(600, len(remaining))
            if cut < len(remaining):
                space = remaining.rfind(" ", 0, cut)
                if space > 350:
                    cut = space
            part, remaining = remaining[:cut], remaining[cut:].lstrip()
            current.append({**b, "text": part})
            if remaining:
                groups.append(current)
                current = []
        last_slide = slide
    if current:
        groups.append(current)
    for index, group in enumerate(groups, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F7F6F1")
        heading = group[0]["text"] if group[0]["type"] == "heading" and len(group[0]["text"]) < 110 else title
        body = group[1:] if heading == group[0]["text"] else group
        for x, y, w, h, content, size, color, is_bold in [(0.8, 0.65, 11.7, 1.1, heading, 30, "244D42", True), (0.8, 2.0, 11.7, 4.6, "\n\n".join(b["text"] for b in body), 19, "222827", False)]:
            tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0)
            tf.text = content
            for p in tf.paragraphs:
                p.font.size = Pt(size)
                p.font.name = "DejaVu Sans"
                p.font.bold = is_bold
                p.font.color.rgb = RGBColor.from_string(color)
                p.space_after = Pt(8)
                if is_rtl(p.text):
                    p.alignment = PP_ALIGN.RIGHT
        footer = slide.shapes.add_textbox(Inches(12), Inches(7), Inches(0.5), Inches(0.2))
        footer.text = str(index)
        footer.text_frame.paragraphs[0].font.size = Pt(10)
    prs.core_properties.title = title
    prs.core_properties.author = "" if clean_metadata else "txtzi"
    prs.core_properties.last_modified_by = ""
    prs.core_properties.comments = ""
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


EXPORTERS = {"docx": (export_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), "pdf": (export_pdf, "application/pdf"), "pptx": (export_pptx, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
